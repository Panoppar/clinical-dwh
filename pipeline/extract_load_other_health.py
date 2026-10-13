"""
Loads the 15 other_health_and_population_tables bronze tables + the
_ingestion_manifest. The most heterogeneous source in this project: a mix
of CSV/XLSX, two files secretly cp874/TIS-620 encoded rather than UTF-8,
one xlsx with a deterministic double-encoding (mojibake) bug, a 228-column
wide table, and a 4-row flattened spreadsheet header. Each table gets its
own small loader function since the shapes don't share enough structure
to generalize further than the common helpers below.
"""

import csv
import io
import json
import os
import re
from datetime import datetime, timezone

import openpyxl
import pymysql

EXTERNAL_SOURCES_ROOT = "/app/external-sources/other_health"
SOURCE_SYSTEM = "other_health_and_population_tables"
BATCH_SIZE = 2000

_manifest_rows = []


def _connect():
    return pymysql.connect(
        host=os.environ["DWH_DB_HOST"],
        port=int(os.environ["DWH_DB_PORT"]),
        user=os.environ["DWH_DB_USER"],
        password=os.environ["DWH_DB_PASSWORD"],
        database="bronze",
        charset="utf8mb4",
        cursorclass=pymysql.cursors.Cursor,
        connect_timeout=10,
    )


def _read_csv_rows(path, skip_rows=1):
    """Reads a CSV, auto-detecting utf-8-sig vs cp874/TIS-620 (two files in
    this source are cp874 despite the .csv extension -- confirmed by
    inspection, not assumed)."""
    with open(path, "rb") as f:
        raw = f.read()
    for enc in ("utf-8-sig", "cp874"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError(f"cannot decode {path!r} as utf-8-sig or cp874")
    reader = csv.reader(io.StringIO(text))
    all_rows = list(reader)
    header = all_rows[:skip_rows]
    data_rows = [r for r in all_rows[skip_rows:] if r and any(r)]
    return header, data_rows


def _insert_batch(cur, insert_sql, rows):
    row_count = 0
    batch = []
    for row in rows:
        batch.append(row)
        if len(batch) >= BATCH_SIZE:
            cur.executemany(insert_sql, batch)
            row_count += len(batch)
            batch.clear()
    if batch:
        cur.executemany(insert_sql, batch)
        row_count += len(batch)
    return row_count


def _record_manifest(batch_id, source_file_name, source_folder, row_count, column_count, status="SUCCESS", error=None):
    _manifest_rows.append((batch_id, source_file_name, source_folder, None, None, None,
                            row_count, column_count, status, error))


def _pad(row, n):
    return (list(row) + [None] * n)[:n]


# ----------------------------------------------------------------------------
def load_sick_298group_long(conn, batch_id) -> int:
    folder = "01_โรคเฝ้าระวังและผู้ป่วย/298กลุ่มโรค_รายจังหวัด_แยกเพศ_ยาว"
    files = [
        "10อันดับกลุ่มโรคผู้ป่วย(298โรค)_สมุทรปราการ_2567.csv",
        "10อันดับกลุ่มโรคผู้ป่วย(298โรค)_สมุทรปราการ_2568.csv",
    ]
    insert_sql = (
        "INSERT INTO sick_298group_long "
        "(source_file_name, source_folder, ingestion_batch_id, source_system, source_row_number, "
        "`ปี`, `จังหวัด`, `ชื่อกลุ่ม`, `เพศ`, `จำนวน`) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)"
    )
    total = 0
    with conn.cursor() as cur:
        for filename in files:
            path = os.path.join(EXTERNAL_SOURCES_ROOT, folder, filename)
            _header, data_rows = _read_csv_rows(path, skip_rows=1)
            rows = [(filename, folder, batch_id, SOURCE_SYSTEM, i, *_pad(r, 5))
                    for i, r in enumerate(data_rows, start=1)]
            n = _insert_batch(cur, insert_sql, rows)
            _record_manifest(batch_id, filename, folder, n, 5)
            total += n
    conn.commit()
    print(f"[other_health] sick_298group_long: loaded {total} rows")
    return total


# ----------------------------------------------------------------------------
WIDE_FILENAME_RE = re.compile(r"^(\d+)_.+_([^_]+)_(\d{4})\.xlsx$")


def load_sick_298group_wide(conn, batch_id) -> int:
    folder = "01_โรคเฝ้าระวังและผู้ป่วย/298กลุ่มโรค_รายจังหวัด_แยกเพศ_กว้าง"
    folder_path = os.path.join(EXTERNAL_SOURCES_ROOT, folder)
    insert_sql = (
        "INSERT INTO sick_298group_wide "
        "(source_file_name, source_folder, ingestion_batch_id, source_system, source_row_number, "
        "file_province, file_year, `ชื่อกลุ่ม(298โรค)`, `ชาย`, `หญิง`, `รวม`) "
        "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)"
    )
    total = 0
    with conn.cursor() as cur:
        for filename in sorted(os.listdir(folder_path)):
            if not filename.endswith(".xlsx"):
                continue
            m = WIDE_FILENAME_RE.match(filename)
            if not m:
                raise ValueError(f"cannot parse {filename!r}")
            _no, province, year = m.groups()
            wb = openpyxl.load_workbook(os.path.join(folder_path, filename), read_only=True, data_only=True)
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            next(rows_iter, None)  # header
            rows = []
            for i, row in enumerate(rows_iter, start=1):
                if row is None or not any(row):
                    continue
                rows.append((filename, folder, batch_id, SOURCE_SYSTEM, i, province, year, *_pad(row, 4)))
            wb.close()
            n = _insert_batch(cur, insert_sql, rows)
            _record_manifest(batch_id, filename, folder, n, 4)
            total += n
    conn.commit()
    print(f"[other_health] sick_298group_wide: loaded {total} rows")
    return total


# ----------------------------------------------------------------------------
def load_surveillance_top_diseases_2566(conn, batch_id) -> int:
    folder = "01_โรคเฝ้าระวังและผู้ป่วย/10อันดับโรคเฝ้าระวัง_2566"
    filename = "10อันดับโรคเฝ้าระวังทางระบาดวิทยา_2566.csv"
    path = os.path.join(EXTERNAL_SOURCES_ROOT, folder, filename)
    _header, data_rows = _read_csv_rows(path, skip_rows=1)
    insert_sql = (
        "INSERT INTO surveillance_top_diseases_2566 "
        "(source_file_name, source_folder, ingestion_batch_id, source_system, source_row_number, "
        "`ปี`, `โรค`, `จำนวน`, `หน่วย`) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)"
    )
    rows = [(filename, folder, batch_id, SOURCE_SYSTEM, i, *_pad(r, 4)) for i, r in enumerate(data_rows, start=1)]
    with conn.cursor() as cur:
        n = _insert_batch(cur, insert_sql, rows)
    conn.commit()
    _record_manifest(batch_id, filename, folder, n, 4)
    print(f"[other_health] surveillance_top_diseases_2566: loaded {n} rows")
    return n


def load_surveillance_top_diseases_2567(conn, batch_id) -> int:
    folder = "01_โรคเฝ้าระวังและผู้ป่วย/10อันดับโรคเฝ้าระวัง_2567"
    filename = "10อันดับโรคเฝ้าระวังทางระบาดวิทยา_2567.csv"
    path = os.path.join(EXTERNAL_SOURCES_ROOT, folder, filename)
    _header, data_rows = _read_csv_rows(path, skip_rows=1)
    insert_sql = (
        "INSERT INTO surveillance_top_diseases_2567 "
        "(source_file_name, source_folder, ingestion_batch_id, source_system, source_row_number, "
        "`ปี`, `ชื่อโรค`, `จำนวน`) VALUES (%s,%s,%s,%s,%s,%s,%s,%s)"
    )
    rows = [(filename, folder, batch_id, SOURCE_SYSTEM, i, *_pad(r, 3)) for i, r in enumerate(data_rows, start=1)]
    with conn.cursor() as cur:
        n = _insert_batch(cur, insert_sql, rows)
    conn.commit()
    _record_manifest(batch_id, filename, folder, n, 3)
    print(f"[other_health] surveillance_top_diseases_2567: loaded {n} rows")
    return n


def load_compare_cases_by_disease(conn, batch_id) -> int:
    folder = "01_โรคเฝ้าระวังและผู้ป่วย/เปรียบเทียบผู้ป่วยรายโรค_2567-2568"
    filename = "จำนวนผู้ป่วยแยกตามโรค_เปรียบเทียบปี2567-2568.csv"
    path = os.path.join(EXTERNAL_SOURCES_ROOT, folder, filename)
    _header, data_rows = _read_csv_rows(path, skip_rows=1)
    insert_sql = (
        "INSERT INTO compare_cases_by_disease "
        "(source_file_name, source_folder, ingestion_batch_id, source_system, source_row_number, "
        "`โรค`, `ปี2567`, `ปี2568`, `ผลต่าง`) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)"
    )
    rows = [(filename, folder, batch_id, SOURCE_SYSTEM, i, *_pad(r, 4)) for i, r in enumerate(data_rows, start=1)]
    with conn.cursor() as cur:
        n = _insert_batch(cur, insert_sql, rows)
    conn.commit()
    _record_manifest(batch_id, filename, folder, n, 4)
    print(f"[other_health] compare_cases_by_disease: loaded {n} rows")
    return n


def load_compare_cases_by_subdistrict(conn, batch_id) -> int:
    folder = "01_โรคเฝ้าระวังและผู้ป่วย/เปรียบเทียบผู้ป่วยรายเขตแขวง_2567-2568"
    filename = "จำนวนผู้ป่วยรายเขต-แขวง_เปรียบเทียบปี2567-2568.csv"
    path = os.path.join(EXTERNAL_SOURCES_ROOT, folder, filename)
    _header, data_rows = _read_csv_rows(path, skip_rows=1)
    insert_sql = (
        "INSERT INTO compare_cases_by_subdistrict "
        "(source_file_name, source_folder, ingestion_batch_id, source_system, source_row_number, "
        "`เขต`, `แขวง`, `ปี2567`, `ปี2568`, `ผลต่าง`) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)"
    )
    rows = [(filename, folder, batch_id, SOURCE_SYSTEM, i, *_pad(r, 5)) for i, r in enumerate(data_rows, start=1)]
    with conn.cursor() as cur:
        n = _insert_batch(cur, insert_sql, rows)
    conn.commit()
    _record_manifest(batch_id, filename, folder, n, 5)
    print(f"[other_health] compare_cases_by_subdistrict: loaded {n} rows")
    return n


# ----------------------------------------------------------------------------
def load_ncd_mortality_multisheet(conn, batch_id) -> int:
    folder = "02_สาเหตุการตายและโรคNCD/NCDรายจังหวัด_multi-sheet"
    filename = "044_จำนวนและอัตราตายด้วย5โรคNCDรายจังหวัด_2562-2566.xlsx"
    path = os.path.join(EXTERNAL_SOURCES_ROOT, folder, filename)
    insert_sql = (
        "INSERT INTO ncd_mortality_multisheet "
        "(source_file_name, source_folder, source_sheet_name, ingestion_batch_id, source_system, source_row_number, "
        "`เขตสุขภาพ`, `จังหวัด`, "
        "`ปี2562_จำนวน`, `ปี2562_อัตรา`, `ปี2562_ประชากรกลางปี`, "
        "`ปี2563_จำนวน`, `ปี2563_อัตรา`, `ปี2563_ประชากรกลางปี`, "
        "`ปี2564_จำนวน`, `ปี2564_อัตรา`, `ปี2564_ประชากรกลางปี`, "
        "`ปี2565_จำนวน`, `ปี2565_อัตรา`, `ปี2565_ประชากรกลางปี`, "
        "`ปี2566_จำนวน`, `ปี2566_อัตรา`, `ปี2566_ประชากรกลางปี`) "
        "VALUES (" + ",".join(["%s"] * 23) + ")"
    )
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    total = 0
    with conn.cursor() as cur:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows_iter = ws.iter_rows(values_only=True)
            for _ in range(5):  # title, blank, blank, year-group header, sub-label header
                next(rows_iter, None)
            rows = []
            for i, row in enumerate(rows_iter, start=1):
                if row is None or not any(row):
                    continue
                # Below the real data table (ending in a legitimate "รวม"
                # nationwide total row) the sheet carries loose footer/
                # attribution text sitting only in column 0 -- confirmed by
                # inspection, e.g. a 'หมายเหตุ : ...' note with every other
                # column None. Real rows (incl. the total row) always carry
                # a จังหวัด value in column 1; footnotes don't.
                if row[1] is None:
                    continue
                rows.append((filename, folder, sheet_name, batch_id, SOURCE_SYSTEM, i, *_pad(row, 17)))
            n = _insert_batch(cur, insert_sql, rows)
            _record_manifest(batch_id, filename, folder, n, 17, status=f"SUCCESS (sheet {sheet_name})")
            total += n
    wb.close()
    conn.commit()
    print(f"[other_health] ncd_mortality_multisheet: loaded {total} rows")
    return total


# ----------------------------------------------------------------------------
def _fix_mojibake(value):
    """Reverses the cp1252->cp874 double-encoding bug confirmed in this
    source's cause_of_death xlsx twin (e.g. 'Ë¹èÇÂ' round-trips to
    'หน่วย' exactly). Deterministic and lossless, so this is corrected at
    read time rather than landed as mojibake -- same treatment as BOM
    stripping elsewhere in this project, not a "fix" to source data values."""
    if not isinstance(value, str):
        return value
    try:
        return value.encode("cp1252").decode("cp874")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return value  # not mojibake (e.g. already-clean ASCII/Thai) -- leave as-is


def load_cause_of_death_top10_with_rate(conn, batch_id) -> int:
    folder = "02_สาเหตุการตายและโรคNCD/สาเหตุการตาย10อันดับ_มีอัตราตาย"
    insert_sql = (
        "INSERT INTO cause_of_death_top10_with_rate "
        "(source_file_name, source_folder, ingestion_batch_id, source_system, source_row_number, "
        "`ปี`, `จังหวัด`, `สาเหตุการตาย`, `เพศชาย`, `เพศหญิง`, `รวม`, `อัตราการตายต่อประชากร`, `หน่วย`, `ที่มา`) "
        "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)"
    )
    total = 0
    with conn.cursor() as cur:
        # CSV twin (cp874-encoded, confirmed by inspection)
        csv_filename = "สาเหตุการตาย10อันดับ(อัตราต่อประชากร)_สมุทรปราการ_2568.csv"
        path = os.path.join(EXTERNAL_SOURCES_ROOT, folder, csv_filename)
        _header, data_rows = _read_csv_rows(path, skip_rows=1)
        rows = [(csv_filename, folder, batch_id, SOURCE_SYSTEM, i, *_pad(r, 9))
                for i, r in enumerate(data_rows, start=1)]
        n = _insert_batch(cur, insert_sql, rows)
        _record_manifest(batch_id, csv_filename, folder, n, 9)
        total += n

        # XLSX twin (mojibake, corrected at read time -- see _fix_mojibake)
        xlsx_filename = "สาเหตุการตาย10อันดับ(อัตราต่อประชากร)_สมุทรปราการ_2568.xlsx"
        wb = openpyxl.load_workbook(os.path.join(EXTERNAL_SOURCES_ROOT, folder, xlsx_filename),
                                     read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        next(rows_iter, None)  # header
        rows = []
        for i, row in enumerate(rows_iter, start=1):
            if row is None or not any(row):
                continue
            fixed = tuple(_fix_mojibake(v) for v in row)
            rows.append((xlsx_filename, folder, batch_id, SOURCE_SYSTEM, i, *_pad(fixed, 9)))
        wb.close()
        n = _insert_batch(cur, insert_sql, rows)
        _record_manifest(batch_id, xlsx_filename, folder, n, 9)
        total += n
    conn.commit()
    print(f"[other_health] cause_of_death_top10_with_rate: loaded {total} rows")
    return total


def load_cause_of_death_top10_no_rate(conn, batch_id) -> int:
    folder = "02_สาเหตุการตายและโรคNCD/สาเหตุการตาย10อันดับ_ไม่มีอัตราตาย"
    filename = "สาเหตุการตาย10อันดับ_สมุทรปราการ_2568.csv"
    path = os.path.join(EXTERNAL_SOURCES_ROOT, folder, filename)
    _header, data_rows = _read_csv_rows(path, skip_rows=1)
    insert_sql = (
        "INSERT INTO cause_of_death_top10_no_rate "
        "(source_file_name, source_folder, ingestion_batch_id, source_system, source_row_number, "
        "`ปี`, `จังหวัด`, `สาเหตุการตาย`, `ชาย`, `หญิง`, `รวม`, `หน่วย`, `ที่มา`) "
        "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)"
    )
    rows = [(filename, folder, batch_id, SOURCE_SYSTEM, i, *_pad(r, 8)) for i, r in enumerate(data_rows, start=1)]
    with conn.cursor() as cur:
        n = _insert_batch(cur, insert_sql, rows)
    conn.commit()
    _record_manifest(batch_id, filename, folder, n, 8)
    print(f"[other_health] cause_of_death_top10_no_rate: loaded {n} rows")
    return n


# ----------------------------------------------------------------------------
def _strip_bom(value):
    if isinstance(value, str):
        return value.replace("﻿", "")
    return value


def load_population_by_age_subdistrict_bkk(conn, batch_id) -> int:
    """9 identifier columns land as named columns; the remaining 219
    age/sex/summary columns collapse into one JSON payload keyed by their
    real header text -- MySQL rejects 219 individual TEXT columns outright
    (confirmed by direct testing: "Row size too large" even with zero other
    columns present), so this mirrors the HDC_manual_exports JSON-payload
    pattern already used elsewhere in this project for the same reason."""
    folder = "03_ประชากร/ประชากรรายอายุ_รายเขตกทม"
    files = [
        "ประชากรรายอายุรายเขต_กรุงเทพมหานคร_ธ.ค.2567.xlsx",
        "ประชากรรายอายุรายเขต_กรุงเทพมหานคร_ธ.ค.2568.xlsx",
    ]
    n_id_cols = 9  # ปีเดือน .. ชื่อหมู่บ้าน
    insert_sql = (
        "INSERT INTO population_by_age_subdistrict_bkk "
        "(source_file_name, source_folder, ingestion_batch_id, source_system, source_row_number, "
        "`ปีเดือน`, `รหัสจังหวัด`, `ชื่อจังหวัด`, `รหัสสำนักทะเบียน`, `ชื่อสำนักทะเบียน`, "
        "`รหัสตำบล`, `ชื่อตำบล`, `รหัสหมู่บ้าน`, `ชื่อหมู่บ้าน`, age_sex_data) "
        "VALUES (" + ",".join(["%s"] * (5 + n_id_cols + 1)) + ")"
    )

    total = 0
    with conn.cursor() as cur:
        for filename in files:
            path = os.path.join(EXTERNAL_SOURCES_ROOT, folder, filename)
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            header = [_strip_bom(h) for h in next(rows_iter)]
            age_sex_cols = header[n_id_cols:]
            assert len(age_sex_cols) == 219, f"expected 219 age/sex columns, found {len(age_sex_cols)}"

            rows = []
            for i, row in enumerate(rows_iter, start=1):
                if row is None or not any(row):
                    continue
                row = [_strip_bom(v) for v in row]
                id_values = _pad(row[:n_id_cols], n_id_cols)
                age_sex_values = (list(row[n_id_cols:]) + [None] * 219)[:219]
                payload = json.dumps(
                    {col: (None if v is None else str(v)) for col, v in zip(age_sex_cols, age_sex_values)},
                    ensure_ascii=False,
                )
                rows.append((filename, folder, batch_id, SOURCE_SYSTEM, i, *id_values, payload))
            wb.close()
            n = _insert_batch(cur, insert_sql, rows)
            _record_manifest(batch_id, filename, folder, n, 228)
            total += n
    conn.commit()
    print(f"[other_health] population_by_age_subdistrict_bkk: loaded {total} rows")
    return total


def load_population_by_tambon_national(conn, batch_id) -> int:
    folder = "03_ประชากร/ประชากรรายตำบล_ทั่วประเทศ"
    filename = "ประชากรรายตำบลทั่วประเทศ_ธ.ค.2568.xlsx"
    path = os.path.join(EXTERNAL_SOURCES_ROOT, folder, filename)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    next(rows_iter, None)  # header
    insert_sql = (
        "INSERT INTO population_by_tambon_national "
        "(source_file_name, source_folder, ingestion_batch_id, source_system, source_row_number, "
        "`YYMM`, `CC-CODE`, `CC-DESC`, `RCODE-CODE`, `RCODE-DESC`, `CCAATT-CODE`, `CCAATT-DESC`, "
        "`CCAATTMM-CODE`, `CCAATTMM-DESC`, `MALE`, `FEMALE`, `TOTAL`, `HOUSE`) "
        "VALUES (" + ",".join(["%s"] * 18) + ")"
    )
    rows = []
    for i, row in enumerate(rows_iter, start=1):
        if row is None or not any(row):
            continue
        # This file is triplicated internally (confirmed by inspection: the
        # header row "ปีเดือน,รหัสจังหวัด,..." repeats verbatim at rows 11387
        # and 22774, exact multiples of the single-copy length) -- not
        # flagged in this source's own DDL, found independently. Same
        # underlying anomaly as the DOPA source's known triplicated
        # snapshot, delivered differently (inline repeated header here vs.
        # 3 separate <table> tags there). No duplicate-flag column exists
        # in this table's schema, so per bronze principle (no dedup) all 3
        # copies of real data are kept; only the 2 header-repeat lines
        # themselves are skipped, since they're not data at all.
        if row[0] == "ปีเดือน":
            continue
        row = [_strip_bom(v) for v in row]
        rows.append((filename, folder, batch_id, SOURCE_SYSTEM, i, *_pad(row, 13)))
    wb.close()
    with conn.cursor() as cur:
        n = _insert_batch(cur, insert_sql, rows)
    conn.commit()
    _record_manifest(batch_id, filename, folder, n, 13)
    print(f"[other_health] population_by_tambon_national: loaded {n} rows")
    return n


# ----------------------------------------------------------------------------
def load_elderly_dependency_by_district(conn, batch_id) -> int:
    folder = "04_ผู้สูงอายุ"
    files = [
        "ผู้สูงอายุติดเตียงติดบ้านติดสังคมรายอำเภอ_สมุทรปราการ_2566.csv",
        "ผู้สูงอายุติดเตียงติดบ้านติดสังคมรายอำเภอ_สมุทรปราการ_2568.csv",
    ]
    insert_sql = (
        "INSERT INTO elderly_dependency_by_district "
        "(source_file_name, source_folder, ingestion_batch_id, source_system, source_row_number, "
        "`ปีงบประมาณ`, `อำเภอ`, `ผู้สูงอายุติดเตียง`, `ผู้สูงอายุติดบ้าน`, `ผู้สูงอายุติดสังคม`, `หน่วย`, `ที่มา`) "
        "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)"
    )
    total = 0
    with conn.cursor() as cur:
        for filename in files:
            path = os.path.join(EXTERNAL_SOURCES_ROOT, folder, filename)
            _header, data_rows = _read_csv_rows(path, skip_rows=1)
            rows = [(filename, folder, batch_id, SOURCE_SYSTEM, i, *_pad(r, 7))
                    for i, r in enumerate(data_rows, start=1)]
            n = _insert_batch(cur, insert_sql, rows)
            _record_manifest(batch_id, filename, folder, n, 7)
            total += n
    conn.commit()
    print(f"[other_health] elderly_dependency_by_district: loaded {total} rows")
    return total


# ----------------------------------------------------------------------------
KPI_FILENAME_RE = re.compile(r"^ตัวชี้วัดสาธารณสุขราย(จังหวัด|อำเภอ)_.*?(?:_ชุด(\d))?\.csv$")


def load_health_kpi_export(conn, batch_id) -> int:
    folder = "05_ตัวชี้วัดสาธารณสุข"
    folder_path = os.path.join(EXTERNAL_SOURCES_ROOT, folder)
    insert_sql = (
        "INSERT INTO health_kpi_export "
        "(source_file_name, source_folder, ingestion_batch_id, source_system, source_row_number, "
        "file_geo_level, file_set_number, "
        "`a_name`, `target`, `result`, `F3`, `result1`, `result2`, `result3`, `result4`, `result5`, "
        "`result6`, `result7`, `result8`, `result9`, `result10`, `result11`, `result12`) "
        "VALUES (" + ",".join(["%s"] * 23) + ")"
    )
    total = 0
    with conn.cursor() as cur:
        for filename in sorted(os.listdir(folder_path)):
            if not filename.endswith(".csv"):
                continue
            m = KPI_FILENAME_RE.match(filename)
            if not m:
                raise ValueError(f"cannot parse {filename!r}")
            geo_level_th, set_number = m.groups()
            geo_level = "province" if geo_level_th == "จังหวัด" else "district"

            path = os.path.join(folder_path, filename)
            _header, data_rows = _read_csv_rows(path, skip_rows=1)
            rows = [(filename, folder, batch_id, SOURCE_SYSTEM, i, geo_level, set_number, *_pad(r, 16))
                    for i, r in enumerate(data_rows, start=1)]
            n = _insert_batch(cur, insert_sql, rows)
            _record_manifest(batch_id, filename, folder, n, 16)
            total += n
    conn.commit()
    print(f"[other_health] health_kpi_export: loaded {total} rows")
    return total


# ----------------------------------------------------------------------------
def load_community_count_by_district_bkk(conn, batch_id) -> int:
    folder = "06_ชุมชนและแรงงานต่างด้าว/จำนวนชุมชนรายเขต"
    filename = "จำนวนชุมชนแยกตามประเภทรายเขต_กรุงเทพมหานคร_2569.xlsx"
    path = os.path.join(EXTERNAL_SOURCES_ROOT, folder, filename)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    next(rows_iter, None)  # header (row0 is the real header here, no title row)
    insert_sql = (
        "INSERT INTO community_count_by_district_bkk "
        "(source_file_name, source_folder, ingestion_batch_id, source_system, source_row_number, "
        "`สำนักงานเขต`, `ชุมชนแออัด`, `ชุมชนชานเมือง`, `ชุมชนอาคารสูง`, `ชุมชนเมือง`, `ชุมชนรูปแบบพิเศษ`) "
        "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)"
    )
    rows = []
    for i, row in enumerate(rows_iter, start=1):
        if row is None or not any(row):
            continue
        rows.append((filename, folder, batch_id, SOURCE_SYSTEM, i, *_pad(row, 6)))
    wb.close()
    with conn.cursor() as cur:
        n = _insert_batch(cur, insert_sql, rows)
    conn.commit()
    _record_manifest(batch_id, filename, folder, n, 6)
    print(f"[other_health] community_count_by_district_bkk: loaded {n} rows")
    return n


def load_foreign_labor_permit_by_area_bkk(conn, batch_id) -> int:
    folder = "06_ชุมชนและแรงงานต่างด้าว/แรงงานต่างด้าวรายอนุญาต"
    filename = "จำนวนแรงงานต่างด้าวที่ได้รับอนุญาตทำงาน_กรุงเทพมหานคร_2568.csv"
    path = os.path.join(EXTERNAL_SOURCES_ROOT, folder, filename)
    # 4-row header: title, group labels (with merge bleed), blank, sub-labels
    _header, data_rows = _read_csv_rows(path, skip_rows=4)
    insert_sql = (
        "INSERT INTO foreign_labor_permit_by_area_bkk "
        "(source_file_name, source_folder, ingestion_batch_id, source_system, source_row_number, "
        "`พื้นที่`, `เขตที่ครอบคลุม`, "
        "`ทั่วไป_รวม`, `ทั่วไป_ชาย`, `ทั่วไป_หญิง`, "
        "`นำเข้าตามMOU_รวม`, `นำเข้าตามMOU_ชาย`, `นำเข้าตามMOU_หญิง`, "
        "`ส่งเสริมการลงทุน_รวม`, `ส่งเสริมการลงทุน_ชาย`, `ส่งเสริมการลงทุน_หญิง`, "
        "`ชนกลุ่มน้อย_รวม`, `ชนกลุ่มน้อย_ชาย`, `ชนกลุ่มน้อย_หญิง`, "
        "`ต่ออายุ_รวม`, `ต่ออายุ_ชาย`, `ต่ออายุ_หญิง`, "
        "`จดทะเบียนสถานะไม่ถูกกฎหมาย_รวม`, `จดทะเบียนสถานะไม่ถูกกฎหมาย_ชาย`, `จดทะเบียนสถานะไม่ถูกกฎหมาย_หญิง`) "
        "VALUES (" + ",".join(["%s"] * 25) + ")"
    )
    rows = [(filename, folder, batch_id, SOURCE_SYSTEM, i, *_pad(r, 20)) for i, r in enumerate(data_rows, start=1)]
    with conn.cursor() as cur:
        n = _insert_batch(cur, insert_sql, rows)
    conn.commit()
    _record_manifest(batch_id, filename, folder, n, 20)
    print(f"[other_health] foreign_labor_permit_by_area_bkk: loaded {n} rows")
    return n


# ----------------------------------------------------------------------------
def flush_manifest(conn, batch_id) -> int:
    insert_sql = (
        "INSERT INTO _ingestion_manifest "
        "(batch_id, source_file_name, source_folder, file_hash, file_size_bytes, sheet_name, "
        "row_count, column_count, status, error_message) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)"
    )
    with conn.cursor() as cur:
        n = _insert_batch(cur, insert_sql, _manifest_rows)
    conn.commit()
    return n


def run() -> None:
    loaded_at = datetime.now(timezone.utc)
    batch_id = f"other_health_load_{loaded_at:%Y%m%dT%H%M%S}"
    conn = _connect()

    total = 0
    total += load_sick_298group_long(conn, batch_id)
    total += load_sick_298group_wide(conn, batch_id)
    total += load_surveillance_top_diseases_2566(conn, batch_id)
    total += load_surveillance_top_diseases_2567(conn, batch_id)
    total += load_compare_cases_by_disease(conn, batch_id)
    total += load_compare_cases_by_subdistrict(conn, batch_id)
    total += load_ncd_mortality_multisheet(conn, batch_id)
    total += load_cause_of_death_top10_with_rate(conn, batch_id)
    total += load_cause_of_death_top10_no_rate(conn, batch_id)
    total += load_population_by_age_subdistrict_bkk(conn, batch_id)
    total += load_population_by_tambon_national(conn, batch_id)
    total += load_elderly_dependency_by_district(conn, batch_id)
    total += load_health_kpi_export(conn, batch_id)
    total += load_community_count_by_district_bkk(conn, batch_id)
    total += load_foreign_labor_permit_by_area_bkk(conn, batch_id)

    flush_manifest(conn, batch_id)
    conn.close()
    print(f"[other_health] total rows loaded: {total}")
